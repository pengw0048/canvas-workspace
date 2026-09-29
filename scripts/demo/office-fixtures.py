"""Writes the Numbers and Keynote source fixtures for the demo: office-fixtures.py <dir>.

Requires openpyxl and python-pptx. Numbers and Keynote import these and save native copies.
"""
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

out = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = "Latency"
rows = [("Region", "p95 before (ms)", "p95 after (ms)"), ("us-east", 42, 17), ("eu-west", 47, 19), ("ap-south", 55, 23), ("sa-east", 51, 21)]
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="0B84F3")
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 18
chart = BarChart()
chart.type = "col"
chart.title = "p95 latency by region"
chart.y_axis.title = "ms"
chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=5), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
chart.width, chart.height = 16, 9
ws.add_chart(chart, "E2")
wb.save(f"{out}/Latency.xlsx")

p = Presentation()
p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
s1 = p.slides.add_slide(p.slide_layouts[0])
s1.shapes.title.text = "Read-through cache"
s1.placeholders[1].text = "Launch review · week 39"
s2 = p.slides.add_slide(p.slide_layouts[5])
s2.shapes.title.text = "p95 latency is down ~60% in every region"
for shape in (s1.shapes.title, s2.shapes.title):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.color.rgb = RGBColor(0x1D, 0x1D, 0x1F)
p.save(f"{out}/Launch review.pptx")
print("wrote", out)
